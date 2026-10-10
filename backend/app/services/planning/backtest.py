"""Geçmiş deneme (backtest): "had the planner run at T0, what would it suggest?"

Hard rules:
  * Excel orders / stock are temporary inputs attached to the backtest run —
    nothing is written to order_lines, stock_movements or products.
  * Only a PlanRun with status=backtest (+ its slots) is written; live
    active/stale plans are never superseded or marked stale.
  * No look-ahead: stats, downtimes, mounted molds and ledger stock are read
    as of T0.
"""

from __future__ import annotations

import io
import re
from datetime import datetime, time, timedelta, timezone
from typing import Any

from sqlalchemy.orm import Session

from app.db.models import OrderImport, OrderLine, PlanRun, Product
from app.services.planning.excel_format import _col_to_index, get_excel_format
from app.services.planning.excel_orders import parse_orders_xlsx
from app.services.planning.excel_stock import parse_stock_xlsx
from app.services.planning.plan_engine import LineIn, compute_plan, gather_context, write_plan_run
from app.services.planning.plan_state import BACKTEST
from app.services.planning.stock import stock_as_of
from app.services.time_windows import DISPLAY_TZ

_T0_RE = re.compile(r"(\d{1,2})[./](\d{1,2})[./](\d{2,4})(?:\s+(\d{1,2})[:.](\d{2}))?")
_CODE_RE = re.compile(r"^[A-Z0-9][A-Z0-9_\-./]{3,}$")


def norm_key(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    return " ".join(str(val).strip().upper().split())


def _first_token(name: str) -> str:
    tok = norm_key(name).split(" ")[0] if name else ""
    return tok if _CODE_RE.match(tok) and any(ch.isdigit() for ch in tok) else ""


class ProductMatcher:
    """Read-only Excel → products matcher (exact name, then product code)."""

    def __init__(self, db: Session):
        self.by_name: dict[str, Product] = {}
        self.by_code: dict[str, Product] = {}
        for p in db.query(Product).all():
            self.by_name.setdefault(norm_key(p.name), p)
            code = _first_token(p.name)
            if code:
                self.by_code.setdefault(code, p)

    def match(self, keys: list[str]) -> tuple[Product | None, str]:
        clean = [norm_key(k) for k in keys if norm_key(k)]
        for k in clean:
            if k in self.by_name:
                return self.by_name[k], "ad"
        for k in clean:
            tok = k.split(" ")[0]
            if tok in self.by_code:
                return self.by_code[tok], "kod"
        return None, "yok"


def detect_t0(ws) -> tuple[datetime | None, str | None]:
    """Find «10.09.26 12:06 TARİHİNDE VERİLER ALINMIŞTIR»-style stamps in the header."""
    for r in range(1, 4):
        for c in range(1, 12):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, datetime) and r == 1 and c <= 2:
                local = v if v.tzinfo else v.replace(tzinfo=DISPLAY_TZ)
                return local.astimezone(timezone.utc), f"{_cell_name(r, c)} hücresi"
            if not isinstance(v, str) or "TAR" not in v.upper():
                continue
            m = _T0_RE.search(v)
            if not m:
                continue
            d, mo, y, hh, mm = m.groups()
            year = int(y) + (2000 if len(y) == 2 else 0)
            local = datetime(year, int(mo), int(d), int(hh or 0), int(mm or 0), tzinfo=DISPLAY_TZ)
            return local.astimezone(timezone.utc), f"{_cell_name(r, c)}: «{v.strip()}»"
    return None, None


def _cell_name(r: int, c: int) -> str:
    from app.services.planning.excel_format import col_letter

    return f"{col_letter(c)}{r}"


def end_of_local_day(dt: datetime) -> datetime:
    d = dt.astimezone(timezone.utc).date()
    return datetime.combine(d, time(23, 59, 59), tzinfo=DISPLAY_TZ).astimezone(timezone.utc)


def inspect_excels(
    db: Session,
    orders_data: bytes,
    stock_data: bytes | None = None,
    *,
    orders_filename: str = "",
    stock_filename: str = "",
) -> dict[str, Any]:
    """Parse with the live excel_format and report mapping/T0/horizon. No DB writes."""
    from openpyxl import load_workbook

    fmt = get_excel_format(db)
    parsed = parse_orders_xlsx(orders_data, db=db)
    wb = load_workbook(io.BytesIO(orders_data), read_only=False, data_only=True)
    ws = wb.active
    t0, t0_source = detect_t0(ws)
    alt_cols = sorted({1, 2, _col_to_index(fmt["orders"].get("product_col") or "B")})
    matcher = ProductMatcher(db)
    warnings: list[str] = list(parsed.get("errors") or [])

    stock_rows: list[dict[str, Any]] = []
    stock_by_raw: dict[str, dict[str, Any]] = {}
    if stock_data:
        sp = parse_stock_xlsx(stock_data, db=db)
        warnings += [f"Stok: {e}" for e in sp.get("errors") or []]
        for r in sp["rows"]:
            prod, how = matcher.match([r["product_name_raw"]])
            raw = norm_key(r["product_name_raw"])
            row = {
                "row": r["row"],
                "product_name_raw": str(r["product_name_raw"]).strip(),
                "qty": int(r["qty"]),
                "product_id": prod.id if prod else None,
                "product_name": prod.name if prod else None,
                "match": how,
                "stock_key": f"p:{prod.id}" if prod else f"raw:{raw}",
            }
            stock_rows.append(row)
            stock_by_raw[raw] = row
            tok = raw.split(" ")[0]
            stock_by_raw.setdefault(tok, row)
        warnings += _stock_header_warning(stock_data, fmt["stock"])

    lines: list[dict[str, Any]] = []
    for i, r in enumerate(parsed["rows"], start=1):
        alts = []
        for c in alt_cols:
            v = ws.cell(row=int(r["row"]), column=c).value
            if v is not None and norm_key(v) and norm_key(v) != norm_key(r["product_name_raw"]):
                alts.append(str(int(v) if isinstance(v, float) and v.is_integer() else v).strip())
        keys = [r["product_name_raw"], *alts]
        prod, how = matcher.match(keys)
        if prod:
            skey = f"p:{prod.id}"
        else:
            skey = None
            for k in keys:
                hit = stock_by_raw.get(norm_key(k)) or stock_by_raw.get(norm_key(k).split(" ")[0])
                if hit:
                    skey = hit["stock_key"]
                    break
            skey = skey or f"raw:{norm_key(r['product_name_raw'])}"
        due = datetime.fromisoformat(r["due_date"]) if r.get("due_date") else None
        lines.append(
            {
                "key": f"x:{i}",
                "row": r["row"],
                "product_name_raw": str(r["product_name_raw"]).strip(),
                "alt_names": alts,
                "product_id": prod.id if prod else None,
                "product_name": prod.name if prod else None,
                "match": how,
                "stock_key": skey,
                "qty": int(r["qty"]),
                "due_date": due.isoformat() if due else None,
                "customer": r.get("customer"),
            }
        )

    dues = [datetime.fromisoformat(ln["due_date"]) for ln in lines if ln["due_date"]]
    if t0 is None and dues:
        t0 = min(dues)
        t0_source = "En erken termin (dosyada veri tarihi bulunamadı)"
        warnings.append("Veri alınma tarihi bulunamadı; başlangıç en erken termin kabul edildi")
    horizon_end = end_of_local_day(max(dues)) if dues else None
    overdue = 0
    if t0 is not None:
        t0_day = t0.astimezone(DISPLAY_TZ).date()
        for ln in lines:
            ln["overdue_at_t0"] = bool(ln["due_date"]) and datetime.fromisoformat(ln["due_date"]).date() < t0_day
            overdue += int(ln["overdue_at_t0"])
    if overdue:
        warnings.append(f"{overdue} sipariş satırının termini veri tarihinden önce (başlangıçta zaten gecikmiş)")

    products: dict[str, dict[str, Any]] = {}
    for ln in lines:
        p = products.setdefault(
            ln["stock_key"],
            {
                "stock_key": ln["stock_key"],
                "product_name_raw": ln["product_name_raw"],
                "alt_names": ln["alt_names"],
                "product_id": ln["product_id"],
                "product_name": ln["product_name"],
                "match": ln["match"],
                "line_count": 0,
                "total_qty": 0,
                "stock": None,
            },
        )
        p["line_count"] += 1
        p["total_qty"] += ln["qty"]
    for s in stock_rows:
        if s["stock_key"] in products:
            products[s["stock_key"]]["stock"] = s["qty"]
    ordered_keys = set(products)
    unused_stock = [s for s in stock_rows if s["stock_key"] not in ordered_keys]
    unmatched = [p for p in products.values() if p["product_id"] is None]
    if unmatched:
        warnings.append(
            f"{len(unmatched)} ürün sistemdeki ürünlerle eşleşmedi — bu siparişler için kalıp/makine önerilemez"
        )

    return {
        "orders_filename": orders_filename,
        "stock_filename": stock_filename,
        "t0": t0.isoformat() if t0 else None,
        "t0_source": t0_source,
        "horizon_end": horizon_end.isoformat() if horizon_end else None,
        "first_due": min(dues).isoformat() if dues else None,
        "line_count": len(lines),
        "product_count": len(products),
        "matched_count": len(products) - len(unmatched),
        "lines": lines,
        "products": sorted(products.values(), key=lambda p: (p["product_id"] is None, p["product_name_raw"])),
        "stock_rows": stock_rows,
        "unused_stock": unused_stock,
        "warnings": warnings,
        "format_used": {"orders": parsed.get("format_used"), "stock": fmt["stock"]},
    }


def _stock_header_warning(data: bytes, cfg: dict[str, Any]) -> list[str]:
    from openpyxl import load_workbook

    start = int(cfg.get("data_start_row") or 1)
    if start <= 1:
        return []
    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    pc, qc = _col_to_index(cfg["product_col"]), _col_to_index(cfg["qty_col"])
    out = []
    for r in range(1, start):
        name = ws.cell(row=r, column=pc).value
        qty = ws.cell(row=r, column=qc).value
        if name is not None and isinstance(qty, (int, float)):
            out.append(
                f"Stok Excel'inde {r}. satır («{name}» = {qty}) veri gibi görünüyor ama ayara göre okunmadı "
                f"(veri başlangıç satırı {start}). Başlık yoksa ayarı 1 yapın."
            )
    return out


def lines_from_report(report: dict[str, Any]) -> list[LineIn]:
    out = []
    for ln in report["lines"]:
        out.append(
            LineIn(
                key=ln["key"],
                product_id=ln["product_id"],
                product_label=ln["product_name"] or ln["product_name_raw"],
                qty=int(ln["qty"]),
                due=datetime.fromisoformat(ln["due_date"]) if ln["due_date"] else None,
                customer=ln.get("customer"),
                stock_key=ln["stock_key"],
            )
        )
    out.sort(key=lambda x: (x.due is None, x.due or datetime.max.replace(tzinfo=timezone.utc), x.key))
    return out


def lines_from_import(db: Session, import_id: int) -> tuple[list[LineIn], dict[str, Any]]:
    """Read-only: an existing order import's lines, treated as open at T0."""
    imp = db.get(OrderImport, int(import_id))
    if not imp:
        raise ValueError("Sipariş listesi bulunamadı")
    rows = (
        db.query(OrderLine)
        .filter(OrderLine.import_id == imp.id, OrderLine.status != "cancelled")
        .order_by(OrderLine.due_date.is_(None), OrderLine.due_date.asc(), OrderLine.id.asc())
        .all()
    )
    lines = [
        LineIn(
            key=f"o:{r.id}",
            order_line_id=r.id,
            product_id=r.product_id,
            product_label=r.product_name_raw,
            qty=int(r.qty),
            due=(r.due_date if r.due_date is None or r.due_date.tzinfo else r.due_date.replace(tzinfo=timezone.utc)),
            customer=r.customer,
        )
        for r in rows
    ]
    dues = [ln.due for ln in lines if ln.due]
    report = {
        "import_id": imp.id,
        "orders_filename": imp.filename,
        "t0": (imp.created_at if imp.created_at.tzinfo else imp.created_at.replace(tzinfo=timezone.utc)).isoformat()
        if imp.created_at
        else None,
        "t0_source": "Sipariş listesinin yüklendiği an",
        "horizon_end": end_of_local_day(max(dues)).isoformat() if dues else None,
        "line_count": len(lines),
        "lines": [
            {
                "key": ln.key,
                "product_name_raw": ln.product_label,
                "product_id": ln.product_id,
                "product_name": ln.product_label,
                "match": "ad" if ln.product_id else "yok",
                "stock_key": ln.skey(),
                "qty": ln.qty,
                "due_date": ln.due.isoformat() if ln.due else None,
                "customer": ln.customer,
            }
            for ln in lines
        ],
        "warnings": [],
    }
    return lines, report


def _parse_dt(val: str | datetime | None) -> datetime | None:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        dt = val
    else:
        dt = datetime.fromisoformat(str(val).replace("Z", "+00:00"))
    return dt if dt.tzinfo else dt.replace(tzinfo=DISPLAY_TZ).astimezone(timezone.utc)


def run_backtest(
    db: Session,
    *,
    orders_data: bytes | None = None,
    stock_data: bytes | None = None,
    import_id: int | None = None,
    as_of: str | datetime | None = None,
    horizon_end: str | datetime | None = None,
    prefs: dict | None = None,
    created_by: str | None = None,
    orders_filename: str = "",
    stock_filename: str = "",
) -> PlanRun:
    if orders_data:
        report = inspect_excels(
            db, orders_data, stock_data, orders_filename=orders_filename, stock_filename=stock_filename
        )
        lines = lines_from_report(report)
    elif import_id is not None:
        lines, report = lines_from_import(db, import_id)
    else:
        raise ValueError("Sipariş Excel'i veya sipariş listesi seçin")

    t0 = _parse_dt(as_of) or _parse_dt(report.get("t0"))
    if t0 is None:
        raise ValueError("Başlangıç zamanı (T0) belirlenemedi — elle girin")
    t1 = _parse_dt(horizon_end) or _parse_dt(report.get("horizon_end")) or (t0 + timedelta(days=45))
    if t1 <= t0:
        raise ValueError("Bitiş, başlangıçtan sonra olmalı")

    stock: dict[str, int] = {}
    stock_conf: dict[str, str] = {}
    if stock_data and report.get("stock_rows") is not None:
        for s in report["stock_rows"]:
            stock[s["stock_key"]] = stock.get(s["stock_key"], 0) + int(s["qty"])
            stock_conf[s["stock_key"]] = "excel"
        report["stock_source"] = "excel"
    else:
        for ln in lines:
            if ln.product_id and ln.skey() not in stock:
                qty, conf = stock_as_of(db, int(ln.product_id), t0)
                stock[ln.skey()] = qty
                stock_conf[ln.skey()] = conf
        report["stock_source"] = "ledger"
        low = sorted({k for k, v in stock_conf.items() if v != "high"})
        if low:
            report.setdefault("warnings", []).append(
                f"{len(low)} ürünün T0 stoğu kayıtlardan kesin çıkarılamadı (Excel stok sayımı yok)"
            )
    report["stock_confidence"] = stock_conf

    ctx = gather_context(db, as_of=t0)
    result = compute_plan(ctx, lines, stock)
    inputs = {k: v for k, v in report.items() if k != "format_used"}
    plan = write_plan_run(
        db,
        result,
        status=BACKTEST,
        created_by=created_by,
        note=f"Geçmiş deneme · {orders_filename or report.get('orders_filename') or 'liste'}"[:256],
        as_of=t0,
        horizon_end=t1,
        prefs=prefs,
        inputs=inputs,
    )
    db.commit()
    db.refresh(plan)
    return plan
