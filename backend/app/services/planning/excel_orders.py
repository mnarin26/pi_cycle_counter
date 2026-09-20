"""Order Excel parser — layout driven by 8080 excel_format settings."""

from __future__ import annotations

import hashlib
import io
from datetime import datetime, timezone
from typing import Any

from app.services.planning.excel_format import _col_to_index, get_excel_format


def _parse_date(val: Any) -> datetime | None:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        dt = val
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
    text = str(val).strip()
    token = text.split()[0] if text else ""
    for fmt in (
        "%Y-%m-%d",
        "%d.%m.%Y",
        "%d/%m/%Y",
        "%Y/%m/%d",
        "%d.%m.%y",
        "%d/%m/%y",
    ):
        try:
            return datetime.strptime(token[:10] if len(token) > 10 else token, fmt).replace(
                tzinfo=timezone.utc
            )
        except ValueError:
            continue
    return None


def _parse_qty(val: Any) -> int:
    if val is None or val == "":
        return 0
    try:
        return int(float(str(val).replace(",", ".").strip()))
    except ValueError:
        return 0


def _cell(ws, row: int, col: int) -> Any:
    return ws.cell(row=row, column=col).value


def _parse_matrix(ws, cfg: dict[str, Any]) -> tuple[list[dict], list[str]]:
    errors: list[str] = []
    parsed: list[dict] = []
    along = cfg.get("dates_along") or "row"

    if along == "column":
        date_col = _col_to_index(cfg["date_col"])
        date_start_row = int(cfg["date_start_row"])
        product_row = int(cfg["product_row"])
        # dates down date_col; products across product_row to the right of date_col
        dates: list[tuple[int, datetime]] = []
        r = date_start_row
        while r < date_start_row + 400:
            d = _parse_date(_cell(ws, r, date_col))
            if d is None and r > date_start_row + 5 and not dates:
                break
            if d is not None:
                dates.append((r, d))
            elif dates and _cell(ws, r, date_col) is None:
                # stop after first empty stretch once we had dates
                if r - dates[-1][0] > 2:
                    break
            r += 1
        products: list[tuple[int, str]] = []
        c = date_col + 1
        while c < date_col + 200:
            name = str(_cell(ws, product_row, c) or "").strip()
            if name:
                products.append((c, name))
            elif products and c > products[-1][0] + 3:
                break
            c += 1
        for prod_c, product in products:
            for date_r, due in dates:
                qty = _parse_qty(_cell(ws, date_r, prod_c))
                if qty <= 0:
                    continue
                parsed.append(
                    {
                        "row": date_r,
                        "product_name_raw": product,
                        "qty": qty,
                        "due_date": due.isoformat(),
                        "customer": None,
                        "note": None,
                    }
                )
        return parsed, errors

    # dates along a row (horizontal)
    product_col = _col_to_index(cfg["product_col"])
    data_start = int(cfg["data_start_row"])
    date_row = int(cfg["date_row"])
    date_start_col = _col_to_index(cfg["date_start_col"])
    customer_col = None
    if cfg.get("customer_col"):
        try:
            customer_col = _col_to_index(cfg["customer_col"])
        except ValueError:
            customer_col = None

    dates: list[tuple[int, datetime]] = []
    c = date_start_col
    empty_streak = 0
    while c < date_start_col + 400:
        d = _parse_date(_cell(ws, date_row, c))
        if d is not None:
            dates.append((c, d))
            empty_streak = 0
        else:
            empty_streak += 1
            if dates and empty_streak > 3:
                break
            if not dates and empty_streak > 10:
                break
        c += 1

    if not dates:
        return [], ["Tarih satırında / kolonunda tarih bulunamadı — Excel format ayarını kontrol edin"]

    max_row = ws.max_row or (data_start + 500)
    for r in range(data_start, min(max_row, data_start + 5000) + 1):
        product = str(_cell(ws, r, product_col) or "").strip()
        if not product:
            continue
        customer = None
        if customer_col:
            customer = str(_cell(ws, r, customer_col) or "").strip() or None
        for col, due in dates:
            qty = _parse_qty(_cell(ws, r, col))
            if qty <= 0:
                continue
            parsed.append(
                {
                    "row": r,
                    "product_name_raw": product,
                    "qty": qty,
                    "due_date": due.isoformat(),
                    "customer": customer,
                    "note": None,
                }
            )
    return parsed, errors


def _parse_table(ws, cfg: dict[str, Any]) -> tuple[list[dict], list[str]]:
    """Classic columns: product / qty / due [/ customer / note]."""
    errors: list[str] = []
    parsed: list[dict] = []
    product_col = _col_to_index(cfg["product_col"])
    qty_col = _col_to_index(cfg.get("qty_col") or "C")
    due_col = _col_to_index(cfg.get("due_col") or "D") if cfg.get("due_col") else None
    customer_col = _col_to_index(cfg["customer_col"]) if cfg.get("customer_col") else None
    note_col = _col_to_index(cfg["note_col"]) if cfg.get("note_col") else None
    start = int(cfg["data_start_row"])
    max_row = ws.max_row or (start + 500)
    for r in range(start, min(max_row, start + 5000) + 1):
        product = str(_cell(ws, r, product_col) or "").strip()
        if not product:
            continue
        qty = _parse_qty(_cell(ws, r, qty_col))
        if qty <= 0:
            errors.append(f"Satır {r}: miktar geçersiz ({product})")
            continue
        due = _parse_date(_cell(ws, r, due_col)) if due_col else None
        customer = str(_cell(ws, r, customer_col) or "").strip() if customer_col else None
        note = str(_cell(ws, r, note_col) or "").strip() if note_col else None
        parsed.append(
            {
                "row": r,
                "product_name_raw": product,
                "qty": qty,
                "due_date": due.isoformat() if due else None,
                "customer": customer or None,
                "note": note or None,
            }
        )
    return parsed, errors


def parse_orders_xlsx(data: bytes, db=None) -> dict[str, Any]:
    """Return preview rows sorted by due date. Does not write DB."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("openpyxl yüklü değil") from e

    file_hash = hashlib.sha256(data).hexdigest()[:32]
    cfg = DEFAULT_CFG_FALLBACK
    if db is not None:
        cfg = get_excel_format(db)["orders"]

    wb = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    ws = wb.active

    if (cfg.get("layout") or "matrix") == "table":
        parsed, errors = _parse_table(ws, cfg)
    else:
        parsed, errors = _parse_matrix(ws, cfg)

    if not parsed and not errors:
        errors = ["Sipariş satırı bulunamadı — Excel format ayarını veya dosyayı kontrol edin"]

    parsed.sort(key=lambda r: (r["due_date"] is None, r["due_date"] or "", r["product_name_raw"]))
    return {
        "file_hash": file_hash,
        "rows": parsed,
        "errors": errors,
        "format_used": cfg,
    }


# Fallback when no db session (tests / CLI)
DEFAULT_CFG_FALLBACK = {
    "layout": "matrix",
    "dates_along": "row",
    "product_col": "B",
    "data_start_row": 3,
    "date_row": 1,
    "date_start_col": "I",
    "date_col": "I",
    "date_start_row": 2,
    "product_row": 2,
    "customer_col": "C",
    "qty_col": "C",
    "due_col": "D",
    "note_col": "E",
}
