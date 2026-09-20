"""Temporary stock Excel parser — column map easy to swap when real format arrives."""

from __future__ import annotations

import hashlib
import io
from typing import Any

STOCK_COLUMNS = {
    "product": ("ürün adı", "urun adi", "ürün", "urun", "product", "product name"),
    "qty": ("stok (adet)", "stok", "adet", "qty", "stock", "miktar"),
    "shipped": ("sevk edilen", "sevk", "shipped"),
}


def _norm(s: Any) -> str:
    return str(s or "").strip().lower().replace("ı", "i").replace("İ", "i")


def _match_col(header: str, aliases: tuple[str, ...]) -> bool:
    h = _norm(header)
    return any(_norm(a) == h or _norm(a) in h for a in aliases)


def _parse_qty(val: Any) -> int:
    if val is None or val == "":
        return 0
    try:
        return int(float(str(val).replace(",", ".").strip()))
    except ValueError:
        return 0


def parse_stock_xlsx(data: bytes) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("openpyxl yüklü değil") from e

    file_hash = hashlib.sha256(data).hexdigest()[:32]
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return {"file_hash": file_hash, "rows": [], "errors": ["Boş dosya"]}

    col_map: dict[str, int] = {}
    for idx, cell in enumerate(header or []):
        for key, aliases in STOCK_COLUMNS.items():
            if key not in col_map and _match_col(str(cell or ""), aliases):
                col_map[key] = idx
    if "product" not in col_map or "qty" not in col_map:
        return {
            "file_hash": file_hash,
            "rows": [],
            "errors": ["Geçici format bekleniyor: Ürün Adı | Stok (adet) [| Sevk Edilen]"],
            "columns_found": [str(c) for c in (header or [])],
        }

    parsed: list[dict] = []
    errors: list[str] = []
    for i, row in enumerate(rows_iter, start=2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        product = str(row[col_map["product"]] or "").strip()
        qty = _parse_qty(row[col_map["qty"]])
        shipped = _parse_qty(row[col_map["shipped"]]) if "shipped" in col_map else None
        if not product:
            errors.append(f"Satır {i}: ürün adı boş")
            continue
        parsed.append(
            {
                "row": i,
                "product_name_raw": product,
                "qty": qty,
                "shipped": shipped,
            }
        )
    parsed.sort(key=lambda r: r["product_name_raw"])
    return {"file_hash": file_hash, "rows": parsed, "errors": errors}
