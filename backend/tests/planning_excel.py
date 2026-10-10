"""Synthetic workbooks with the same layout as the customer's order/stock Excel.

Layout (matches the live excel_format defaults): stamp in A1, delivery dates
from I1 along row 1 with blank weekly-total columns, headers in row 2, data
from row 3; material code in A, description in B, customer in C. Stock is a
header-less code/qty sheet.
"""

from __future__ import annotations

import io

from openpyxl import Workbook

ORDER_DATES = [
    "09.09.26", "10.09.26", "11.09.26", None, "12.09.26", "14.09.26", "15.09.26",
    "16.09.26", "17.09.26", "18.09.26", None, "19.09.26", "21.09.26", "22.09.26",
]

# (code, description, {date_index: qty})
ORDER_ROWS = [
    (5723410100, "YENI_ORTA_MENTESE_TAPASI", {0: 300, 5: 2000, 11: 2500}),
    (4306080600, "KAPI USTU ALIN GR-60CM-BEKO-ISIKSIZ", {9: 40, 12: 30}),
    (6101811700, "NEM SENSORLU MENTESE KAP. GR. KK ARCP1", {7: 1200}),
    (1163990200, "BASTET_LOWER_DOOR_BOT_RIGHT_CAP_BLACK", {11: 33}),
]

STOCK_ROWS = [
    (115979, 100),
    (5723410100, 500),
    (4306080600, 70),
    ("AK0131 ", 52),
    (6101811700, 0),
]


def orders_xlsx(stamp: str = "10.09.26 12:06 TARİHİNDE VERİLER ALINMIŞTIR") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sayfa1"
    ws.cell(row=1, column=1, value=stamp)
    ws.cell(row=1, column=6, value="TESLİM TARİHİ")
    for i, d in enumerate(ORDER_DATES):
        if d:
            ws.cell(row=1, column=9 + i, value=d)
        else:
            ws.cell(row=2, column=9 + i, value="t")
    for c, h in enumerate(["Malzeme", "Malzeme Tanımı", "FİRMA", " Revision", "ASRS"], start=1):
        ws.cell(row=2, column=c, value=h)
    for r, (code, name, qtys) in enumerate(ORDER_ROWS, start=3):
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=112566)
        for idx, q in qtys.items():
            ws.cell(row=r, column=9 + idx, value=q)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def stock_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    for r, (code, qty) in enumerate(STOCK_ROWS, start=1):
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=qty)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
